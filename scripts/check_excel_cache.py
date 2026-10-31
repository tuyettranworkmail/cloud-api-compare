"""Verify native Excel cached results against Python after COM recalculation."""
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from benchmark.io import file_hash, utcnow, write_json


def main():
    path = Path("reports/generated/demo/benchmark_results_demo.xlsx")
    book = load_workbook(path, data_only=True)
    runs = pd.read_csv("results/processed/demo/per_run.csv")
    selected = runs.loc[(runs["scenario"] == "steady") & (runs["load_level"] == "medium") & (runs["endpoint"] == "all") & runs["valid_for_comparison"]]
    checked = []
    for row, model in enumerate(("iaas", "paas", "faas"), 6):
        expected = selected.loc[selected["model"] == model, "p95_ms"].mean()
        actual = book["Overview"].cell(row, 4).value
        assert actual is not None and abs(actual - expected) < 1e-8
        checked.append({"model": model, "excel_mean_run_p95_ms": actual, "python_mean_run_p95_ms": expected})
    statuses = [book["Cost Summary"].cell(row, 12).value for row in range(6, 15)]
    assert statuses == ["OK"] * 9
    formula_errors = []
    for sheet in book:
        for row in sheet:
            for cell in row:
                if cell.data_type == "e":
                    formula_errors.append(f"{sheet.title}!{cell.coordinate}: {cell.value}")
    assert not formula_errors
    sheets = book.sheetnames
    book.close()
    write_json(Path("reports/excel_native_verification.json"), {"verified_at": utcnow(), "workbook_sha256": file_hash(path), "native_excel_recalculated": True, "formula_errors": formula_errors, "cost_reconciliation": statuses, "p95_cross_checks": checked, "sheets": sheets, "rendered_sheet_previews": len(list(Path("reports/excel_previews").glob("*.pdf"))), "data_mode": "DEMO"})
    print("PASS: native Excel cache agrees with Python; 9 cost checks OK; no Excel errors")


if __name__ == "__main__":
    main()
