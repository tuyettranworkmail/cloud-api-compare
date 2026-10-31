"""Portable Excel report exporter (openpyxl fallback explicitly selected by user).

Scientific estimates are snapshots from the Python pipeline. Spreadsheet formulas
summarize those snapshots; this workbook is not a second implementation of SciPy
bootstrap or the occupancy cost model. Change assumptions in YAML and rerun CLI.
"""
from __future__ import annotations

import json
import math
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

from benchmark.io import read_json, write_json
from benchmark.visualization.charts import COLORS, read_table

NAVY, BLUE, AMBER = "142B4A", "2563A6", "FEF3C7"
HEADER = 5
START = 6


def value_for_cell(value):
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False)
    if value is None or (isinstance(value, float) and not math.isfinite(value)) or pd.isna(value):
        return None
    if hasattr(value, "item"):
        value = value.item()
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        # Imported values are source data, not formulas supplied by a remote API.
        return "'" + value
    return value


def export_workbook(analysis: Path, financial: Path, output: Path) -> Path:
    provenance = read_json(analysis / "provenance.json")
    mode = provenance["data_modes"]
    if len(mode) != 1 or mode[0] not in {"DEMO", "REAL"}:
        raise ValueError("Excel report cần một mode DEMO hoặc REAL riêng biệt")
    mode = mode[0]
    label = provenance["label"]
    catalog = read_json(financial / "pricing_catalog.json")
    workload = read_json(financial / "workload_assumptions.json")
    if workload["data_mode"] != mode:
        raise ValueError("Excel cost/performance mode không khớp")
    wb = Workbook()
    cover = wb.active
    cover.title = "Overview"
    wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True)
    wb.properties.title = "Cloud Benchmark • Phúc Hoàng"
    wb.properties.subject = label
    wb.properties.description = "Research report snapshots; recompute scientific and financial models with the Python CLI."
    wb.properties.creator = "Phúc Hoàng • CloudBench"
    schema = {}

    def decorate(ws, title, source, width):
        ws.sheet_view.showGridLines = False
        end = get_column_letter(max(8, width))
        for row, text, fill, color, size in [(1, title, NAVY, "FFFFFF", 19), (2, label, AMBER if mode == "DEMO" else "DBEAFE", "92400E" if mode == "DEMO" else NAVY, 10), (3, str(source), "F1F5F9", "475569", 10)]:
            ws.merge_cells(f"A{row}:{end}{row}")
            cell = ws.cell(row, 1, text)
            cell.fill, cell.font = PatternFill("solid", fgColor=fill), Font(name="Calibri", size=size, bold=row < 3, color=color)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.row_dimensions[row].height = 34 if row == 1 else 30
        ws.freeze_panes = "D6"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = "1:5"
        ws.print_options.horizontalCentered = False
        ws.oddFooter.center.text = mode + " | &A | &P / &N"

    def table_sheet(name: str, data: pd.DataFrame, source: Path | str):
        ws = wb.create_sheet(name)
        decorate(ws, name, source, len(data.columns))
        for index, column in enumerate(data.columns, 1):
            cell = ws.cell(HEADER, index, column)
            cell.fill = PatternFill("solid", fgColor=BLUE)
            cell.font = Font(name="Calibri", size=10, color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(index)].width = 21 if any(x in column for x in ("_ms", "rps", "usd", "count", "ratio")) else 26
            if any(x in column for x in ("method", "source", "warning", "note", "run_id")):
                ws.column_dimensions[get_column_letter(index)].width = 42
        ws.row_dimensions[HEADER].height = 36
        for row_index, row in enumerate(data.itertuples(index=False, name=None), START):
            for column_index, value in enumerate(row, 1):
                cell = ws.cell(row_index, column_index, value_for_cell(value))
                cell.font = Font(name="Calibri", size=10, color=NAVY)
                cell.alignment = Alignment(vertical="center", wrap_text=isinstance(value, str))
                if isinstance(value, (float, int)) and not isinstance(value, bool):
                    cell.number_format = '"$"#,##0.0000;[Red]("$"#,##0.0000);"$"0.0000' if "usd" in data.columns[column_index - 1] else "0.00" if isinstance(value, float) else "#,##0"
                if isinstance(value, str) and value.startswith("http"):
                    cell.hyperlink = value
                    cell.font = Font(name="Calibri", size=10, color=BLUE, underline="single")
            ws.row_dimensions[row_index].height = 32
        if len(data) and len(data.columns):
            ref = f"A{HEADER}:{get_column_letter(len(data.columns))}{HEADER + len(data)}"
            table = Table(displayName="T" + name.replace(" ", "").replace("-", ""), ref=ref)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(table)
        schema[name] = {str(col): get_column_letter(i + 1) for i, col in enumerate(data.columns)}
        return ws

    imported = {}
    for name, filename in [("Per Run", "per_run"), ("Confidence", "comparison"), ("Cold Warm", "cold_warm"), ("Scaling", "scaling"), ("Burst Phases", "burst_phases"), ("Errors", "errors")]:
        data = read_table(analysis, filename)
        imported[name] = data
        table_sheet(name, data, analysis / f"{filename}.csv")
    components = read_table(financial, "cost_breakdown")
    levels = read_table(financial, "cost_levels")
    table_sheet("Cost Breakdown", components, financial / "cost_breakdown.csv")
    # Inputs are imported Python results; totals and reconciliation remain live formulas.
    summary = levels[["level", "model", "monthly_requests", "total_usd", "compute_usd", "shared_data_usd", "without_shared_usd", "cost_per_million_usd", "capacity_status"]].copy()
    summary.rename(columns={"total_usd": "python_total_usd", "compute_usd": "python_compute_usd", "shared_data_usd": "python_shared_usd", "without_shared_usd": "python_without_shared_usd", "cost_per_million_usd": "python_unit_usd"}, inplace=True)
    summary["calculated_total_usd"] = None
    summary["difference_usd"] = None
    summary["check"] = None
    summary["lookup_key"] = None
    cost_ws = table_sheet("Cost Summary", summary, financial / "cost_levels.csv; totals recalculate from Cost Breakdown")
    component_last = HEADER + len(components)
    cc = schema["Cost Breakdown"]
    def cr(column):
        return f"'Cost Breakdown'!${cc[column]}${START}:${cc[column]}${max(START, component_last)}"
    for row in range(START, START + len(summary)):
        cost_ws[f"J{row}"] = f'=SUMIFS({cr("usd")},{cr("level")},A{row},{cr("model")},B{row})'
        cost_ws[f"K{row}"] = f"=J{row}-D{row}"
        cost_ws[f"L{row}"] = f'=IF(ABS(K{row})<0.000001,"OK","CHECK")'
        cost_ws[f"M{row}"] = f'=A{row}&"|"&B{row}'
        for column in ("J", "K"):
            cost_ws[f"{column}{row}"].number_format = '"$"0.000000'
            cost_ws[f"{column}{row}"].font = Font(name="Calibri", size=10, color="008000" if column == "J" else "000000")
    cost_ws.conditional_formatting.add(f"K{START}:K{HEADER+len(summary)}", CellIsRule(operator="notBetween", formula=["-0.000001", "0.000001"], fill=PatternFill("solid", fgColor="FEE2E2")))
    table_sheet("Sensitivity", read_table(financial, "sensitivity"), financial / "sensitivity.csv — frozen Python scenario runs; rerun CLI after edits")
    price_rows = []
    for key, rate in catalog["rates"].items():
        price_rows.append({"key": key, "service": rate["service"], "sku_description": rate["sku"], "unit": rate["unit"], "unit_price_usd": rate["unit_price"], "region": catalog["region"], "retrieved_date": catalog["retrieved_date"], "effective_date": catalog["effective_date"], "source_url": rate["source"], "status": "MISSING_UPDATE_BEFORE_USE" if rate["unit_price"] is None else "PUBLIC_LIST_PRICE_SNAPSHOT"})
    price_ws = table_sheet("Pricing Sources", pd.DataFrame(price_rows), "configs/pricing_catalog.yaml; reports/sources/source_manifest.json")
    price_ws.column_dimensions["C"].width = 42
    for row, rate in enumerate(price_rows, START):
        price_ws.cell(row, 5).number_format = '"$"0.000000000'
        price_ws.cell(row, 5).comment = Comment(rate["source_url"] + "\nRetrieved: " + str(catalog["retrieved_date"]), "Source")
        price_ws.row_dimensions[row].height = 60
        if rate["status"].startswith("MISSING"):
            price_ws.cell(row, 5).fill = PatternFill("solid", fgColor="FEF3C7")

    assumption_rows = []
    def flatten(prefix, value):
        if isinstance(value, dict):
            for key, item in value.items():
                flatten(prefix + "." + key if prefix else key, item)
        else:
            assumption_rows.append({"parameter": prefix, "value": value, "source": "configs/cost_workloads.yaml", "note": "Snapshot; change YAML and rerun cost/export. Not a live Excel occupancy model."})
    flatten("", workload)
    assumption_ws = table_sheet("Assumptions", pd.DataFrame(assumption_rows), financial / "workload_assumptions.json")
    for column, width in {"A": 52, "B": 25, "C": 44, "D": 75}.items():
        assumption_ws.column_dimensions[column].width = width
    for row in range(START, START + len(assumption_rows)):
        assumption_ws.row_dimensions[row].height = 44
    roots = read_json(financial / "break_even.json")
    table_sheet("Break Even", pd.DataFrame([{"left": item["left"], "right": item["right"], "status": item["status"], "roots_requests": item["roots"], "coincident_intervals": item["coincident_intervals"], "discontinuities": item["discontinuities"], "limitation": item["search_resolution"]["limitation"]} for item in roots]), financial / "break_even.json")
    notes = [
        ("Data mode", mode), ("Raw data", provenance["input_root"]), ("Units", "Latency ms; durations s; throughput req/s; money USD/month; ratios already in percent where _pct"),
        ("Statistics", "Per-run metric snapshots; CI from independent-run bootstrap, not pooled request percentiles"),
        ("Raw requests", str(analysis / "requests.csv") + "; original JSONL retained under dataset raw/"),
        ("Outliers and errors", "Retained. Invalid/conflicting records quarantined; see quality.json"),
        ("Cold scope", "Process-first != infrastructure cold. Confirmed needs independent platform evidence"),
        ("Scaling scope", "Minute samples; interval bounds. IaaS fixed VM = N/A"),
        ("Cost scope", "Estimate from public prices + explicit workload assumptions, not invoice"),
        ("Workbook calculations", "Overview and cost reconciliation are Excel formulas. Complex model/CI are frozen Python outputs"),
        ("Refresh", "Run analyze -> cost -> export. Excel fullCalcOnLoad recalculates formula summaries"),
        ("Excel colors", "Blue inputs/selectors; green worksheet references; black calculations; amber missing/DEMO"),
        ("Real measurement", "NOT EXECUTED" if mode == "DEMO" else "Review metadata/validity and invoice scope"),
    ]
    table_sheet("Method and Provenance", pd.DataFrame(notes, columns=["item", "description"]), "README.md; reports/methodology.md; docs/metric_definitions.md")
    wb["Method and Provenance"].column_dimensions["A"].width = 28
    wb["Method and Provenance"].column_dimensions["B"].width = 100
    for row in range(START, START + len(notes)):
        wb["Method and Provenance"].row_dimensions[row].height = 38

    decorate(cover, "CloudBench | Research overview", "Cohort selectors: A4 scenario; B4 load; C4 endpoint. H4 cost level. Only valid independent runs.", 9)
    cover.freeze_panes = "C6"
    for cell, value in {"A4": "steady", "B4": "medium", "C4": "all", "H4": "medium"}.items():
        cover[cell] = value
        cover[cell].font = Font(name="Calibri", size=11, color="0000FF", bold=True)
        cover[cell].fill = PatternFill("solid", fgColor="DBEAFE")
    for cell, choices in (("A4", "steady,burst,idle"), ("B4", "low,medium,high,burst,idle"), ("C4", "all,read,write,cpu"), ("H4", "low,medium,high")):
        dv = DataValidation(type="list", formula1=f'"{choices}"')
        cover.add_data_validation(dv)
        dv.add(cover[cell])
    headings = ["Model", "Valid runs", "Mean run p50 (ms)", "Mean run p95 (ms)", "Mean run p99 (ms)", "Throughput (req/s)", "Error rate (%)", "Monthly cost (USD)", "Capacity status"]
    for column, heading in enumerate(headings, 1):
        cell = cover.cell(HEADER, column, heading)
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cover.column_dimensions[get_column_letter(column)].width = 21 if column < 9 else 34
    cover.row_dimensions[HEADER].height = 40
    per = imported["Per Run"]
    pc = schema["Per Run"]
    last = max(START, HEADER + len(per))
    def pr(key):
        return f"'Per Run'!${pc[key]}${START}:${pc[key]}${last}"
    for row, model in enumerate(("iaas", "paas", "faas"), START):
        cover.cell(row, 1, model.upper()).font = Font(name="Calibri", size=12, color=COLORS[model][1:], bold=True)
        criteria = f'{pr("model")},LOWER(A{row}),{pr("scenario")},$A$4,{pr("load_level")},$B$4,{pr("endpoint")},$C$4,{pr("valid_for_comparison")},TRUE'
        cover.cell(row, 2, "=COUNTIFS(" + criteria + ")")
        for column, metric in enumerate(("p50_ms", "p95_ms", "p99_ms", "achieved_rps", "error_rate_pct"), 3):
            cover.cell(row, column, f'=IFERROR(AVERAGEIFS({pr(metric)},{criteria}),"N/A")')
        end = HEADER + len(summary)
        cover.cell(row, 8, f'=SUMIFS(\'Cost Summary\'!$J${START}:$J${end},\'Cost Summary\'!$A${START}:$A${end},$H$4,\'Cost Summary\'!$B${START}:$B${end},LOWER(A{row}))')
        cover.cell(row, 9, f'=IFERROR(INDEX(\'Cost Summary\'!$I${START}:$I${end},MATCH($H$4&"|"&LOWER(A{row}),\'Cost Summary\'!$M${START}:$M${end},0)),"N/A")')
        for cell in cover[row][:9]:
            cell.fill = PatternFill("solid", fgColor="F1F5F9" if row % 2 == 0 else "FFFFFF")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.data_type == "f":
                cell.font = Font(name="Calibri", size=12, color="008000")
                cell.number_format = '"$"#,##0.00' if cell.column == 8 else "0.00"
        cover.row_dimensions[row].height = 40
    cover.merge_cells("A10:I10")
    cover["A10"] = "INCONCLUSIVE for empirical GCP conclusions in DEMO. Capacity is a separate requirement; latency and cost have different scopes."
    cover["A10"].alignment = Alignment(wrap_text=True, vertical="center")
    cover["A10"].font = Font(name="Calibri", size=11, color="92400E", bold=True)
    cover.row_dimensions[10].height = 34
    chart = BarChart()
    chart.title = "Mean run p95 • selected cohort"
    chart.y_axis.title = "Latency (ms)"
    chart.x_axis.title = None  # Model category labels already communicate the dimension.
    chart.add_data(Reference(cover, min_col=4, min_row=5, max_row=8), titles_from_data=True)
    chart.set_categories(Reference(cover, min_col=1, min_row=6, max_row=8))
    chart.legend = None
    chart.height, chart.width = 11, 28
    chart.style = 10
    chart.y_axis.scaling.min = 0
    chart.y_axis.numFmt = '0.0'
    chart.y_axis.tickLblPos = "nextTo"
    chart.x_axis.tickLblPos = "nextTo"
    chart.x_axis.delete = chart.y_axis.delete = False
    for index, model in enumerate(("iaas", "paas", "faas")):
        point = DataPoint(idx=index)
        point.spPr.solidFill = COLORS[model][1:]
        point.spPr.ln.solidFill = COLORS[model][1:]
        chart.series[0].data_points.append(point)
    cover.add_chart(chart, "A12")
    cover["A37"] = "Bảng nguồn đầy đủ ở Per Run / Confidence; không dùng biểu đồ mean như kiểm định thống kê."
    cover.merge_cells("A37:I37")
    cover["A37"].alignment = Alignment(wrap_text=True)
    cover.print_area = "A1:I39"
    # Add a compact native cost chart with a formula-backed helper table.
    for column, header in enumerate(["Model", "Low USD", "Medium USD", "High USD"], 14):
        cost_ws.cell(HEADER, column, header)
        cost_ws.column_dimensions[get_column_letter(column)].width = 18
    for row, model in enumerate(("iaas", "paas", "faas"), START):
        cost_ws.cell(row, 14, model.upper())
        for column, level in enumerate(("low", "medium", "high"), 15):
            cost_ws.cell(row, column, f'=SUMIFS($J${START}:$J${HEADER+len(summary)},$A${START}:$A${HEADER+len(summary)},"{level}",$B${START}:$B${HEADER+len(summary)},LOWER(N{row}))')
            cost_ws.cell(row, column).number_format = '"$"0.00'
    cost_chart = BarChart()
    cost_chart.title = "Estimated monthly cost"
    cost_chart.y_axis.title = "USD / month"
    cost_chart.x_axis.title = None
    cost_chart.add_data(Reference(cost_ws, min_col=15, max_col=17, min_row=5, max_row=8), titles_from_data=True)
    cost_chart.set_categories(Reference(cost_ws, min_col=14, min_row=6, max_row=8))
    cost_chart.y_axis.scaling.min = 0
    cost_chart.y_axis.numFmt = '"$"0'
    cost_chart.y_axis.tickLblPos = cost_chart.x_axis.tickLblPos = "nextTo"
    cost_chart.x_axis.delete = cost_chart.y_axis.delete = False
    cost_chart.height, cost_chart.width = 10, 25
    cost_chart.legend.position = "b"
    for row, text in ((1, "Cost comparison | formula-linked totals"), (2, label), (3, "Cost Summary and Cost Breakdown; USD/month estimates; capacity separately validated")):
        cost_ws.merge_cells(start_row=row, start_column=14, end_row=row, end_column=27)
        cell = cost_ws.cell(row, 14, text)
        cell.fill = PatternFill("solid", fgColor=NAVY if row == 1 else AMBER if row == 2 else "F1F5F9")
        cell.font = Font(name="Calibri", size=13 if row == 1 else 10, bold=row < 3, color="FFFFFF" if row == 1 else "92400E" if row == 2 else "475569")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    cost_ws.add_chart(cost_chart, "N11")
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    check = load_workbook(output, data_only=False)
    formula_count = sum(cell.data_type == "f" for ws in check for row in ws for cell in row)
    write_json(output.with_suffix(".verification.json"), {"data_mode": mode, "sheets": check.sheetnames, "formulas": formula_count, "native_charts": sum(len(ws._charts) for ws in check), "recalculation": "Excel fullCalcOnLoad; delivery additionally recalculated by native Excel if available", "scientific_model": "Python snapshots; rerun CLI for CI/occupancy/sensitivity changes", "source_analysis": str(analysis), "source_financial": str(financial)})
    check.close()
    return output
