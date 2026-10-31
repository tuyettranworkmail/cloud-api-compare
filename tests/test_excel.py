from pathlib import Path

from openpyxl import load_workbook

from benchmark.visualization.excel import export_workbook


def test_excel_export_has_formula_totals_sources_and_native_charts(tmp_path):
    output = export_workbook(Path("results/processed/demo"), Path("results/financial/demo"), tmp_path / "report.xlsx")
    book = load_workbook(output, data_only=False)
    assert "Overview" in book.sheetnames
    assert "Pricing Sources" in book.sheetnames
    assert book["Overview"]["D6"].data_type == "f"
    assert book["Cost Summary"]["J6"].value.startswith("=SUMIFS")
    assert book["Cost Summary"]["L6"].data_type == "f"
    assert sum(len(sheet._charts) for sheet in book) == 2
    assert "SIMULATED DATA" in book["Overview"]["A2"].value
    assert book["Pricing Sources"]["E6"].comment is not None
    assert book["Per Run"].freeze_panes == "D6"
    assert book.calculation.fullCalcOnLoad
    book.close()
